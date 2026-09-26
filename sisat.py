import io
import json
import random
from datetime import date, datetime
from fractions import Fraction
from html import escape

import xlsxwriter
from flask import request, redirect, session, flash, send_file
from openpyxl import load_workbook

import app as core


SKILLS = {
    'lectura': {
        'title': 'Toma de lectura',
        'components': [
            ('fluidez', 'Fluidez', ('La lectura es fluida', 'La lectura es parcialmente fluida', 'No hay fluidez en la lectura')),
            ('precision', 'Precisión', ('Precisión en la lectura', 'Precisión moderada en la lectura', 'Imprecisión en la lectura')),
            ('palabras_complejas', 'Atención a palabras complejas', ('Atención a palabras complejas', 'Atención en algunas palabras complejas', 'Sin atención a palabras complejas')),
            ('voz', 'Uso de la voz al leer', ('Uso adecuado de la voz al leer', 'Uso inconsistente de la voz al leer', 'Manejo inadecuado de la voz al leer')),
            ('seguridad', 'Seguridad y disposición', ('Seguridad y disposición ante la lectura', 'Seguridad limitada y esfuerzo ante la lectura', 'Inseguridad o indiferencia ante la lectura')),
            ('comprension', 'Comprensión general', ('Comprensión general del texto', 'Comprensión parcial del texto', 'Comprensión deficiente del texto')),
        ],
    },
    'escritura': {
        'title': 'Producción de textos escritos',
        'components': [
            ('legibilidad', 'Legibilidad', ('Es legible', 'Es medianamente legible', 'No se puede leer')),
            ('proposito', 'Propósito comunicativo', ('Cumple con el propósito comunicativo', 'Cumple parcialmente con el propósito comunicativo', 'No cumple con el propósito comunicativo')),
            ('concordancia', 'Relación de palabras y oraciones', ('Relaciona adecuadamente palabras y oraciones', 'Relaciona correctamente algunas palabras u oraciones', 'No relaciona palabras u oraciones')),
            ('vocabulario', 'Diversidad de vocabulario', ('Diversidad de vocabulario', 'Uso limitado de vocabulario', 'Vocabulario limitado o no pertinente')),
            ('puntuacion', 'Signos de puntuación', ('Uso de los signos de puntuación', 'Uso de algunos signos de puntuación', 'No utiliza signos de puntuación')),
            ('ortografia', 'Reglas ortográficas', ('Uso correcto de las reglas ortográficas', 'Uso de algunas reglas ortográficas', 'No respeta las reglas ortográficas')),
        ],
    },
    'calculo': {'title': 'Cálculo mental', 'components': []},
}

LEVEL_COLORS = {'NIVEL ESPERADO': '#DFF2E5', 'EN DESARROLLO': '#FFF0C9', 'REQUIERE APOYO': '#FADDDD'}

RUBRIC_GUIDANCE = {
    'lectura': {
        'fluidez': {
            3: 'Lee palabras, frases y oraciones completas con ritmo y claridad; realiza las pausas indicadas por la puntuación.',
            2: 'Lee con ritmo solo algunas oraciones o párrafos y atiende únicamente algunos signos de puntuación.',
            1: 'Lee de manera monótona e imprecisa y hace pausas constantes que no corresponden a la puntuación.',
        },
        'precision': {
            3: 'Lee correctamente palabras conocidas y desconocidas y articula sin dificultad.',
            2: 'Vacila, sustituye u omite palabras; comete hasta 5% de errores o presenta dificultad con sílabas trabadas.',
            1: 'Comete más de 6% de errores, hace falsos inicios o no logra articular palabras con sílabas trabadas.',
        },
        'palabras_complejas': {
            3: 'Lee cuidadosamente palabras complejas o desconocidas, sin titubeos ni sustituciones.',
            2: 'Se detiene y corrige algunas palabras complejas, o las sustituye por otras similares.',
            1: 'Se equivoca y continúa sin corregir, o evita leer palabras complejas o desconocidas.',
        },
        'voz': {
            3: 'Emplea volumen, entonación y dicción adecuados y atiende signos interrogativos y exclamativos.',
            2: 'El volumen y la entonación son adecuados solo en partes; presenta inconsistencias de expresividad o dicción.',
            1: 'La lectura es monótona, con problemas graves de volumen, entonación o dicción.',
        },
        'seguridad': {
            3: 'Muestra actitud positiva, dominio de la práctica lectora y disfrute al leer.',
            2: 'Presenta tensión o dificultad, aunque logra manejar el momento y continuar.',
            1: 'Muestra contrariedad, nerviosismo que interfiere o apatía ante la lectura.',
        },
        'comprension': {
            3: 'Comunica información específica, reconoce ideas principales, personajes y escenarios, y emite una opinión.',
            2: 'Expone datos generales o algunas ideas y tiene dificultad para emitir una opinión.',
            1: 'No recupera información del texto, no relaciona sus elementos o no logra emitir una opinión.',
        },
    },
    'escritura': {
        'legibilidad': {
            3: 'Separa correctamente las palabras; el trazo, tamaño y organización del texto permiten leerlo.',
            2: 'Presenta algunos errores de separación, trazo o distribución que dificultan parcialmente la lectura.',
            1: 'Los errores de separación, trazo, tamaño u organización impiden leer el texto.',
        },
        'proposito': {
            3: 'Expone, describe, narra o argumenta lo solicitado y corresponde al tipo de texto indicado.',
            2: 'Cumple solo parcialmente la intención o las características del tipo de texto solicitado.',
            1: 'No desarrolla la intención solicitada ni corresponde al tipo de texto requerido.',
        },
        'concordancia': {
            3: 'Relaciona enunciados con nexos pertinentes, mantiene concordancia y usa correctamente los tiempos verbales.',
            2: 'Relaciona algunos enunciados, pero usa pocos nexos y presenta errores de concordancia o conjugación.',
            1: 'No relaciona enunciados, carece de nexos y muestra errores evidentes de concordancia y tiempos verbales.',
        },
        'vocabulario': {
            3: 'Usa vocabulario adecuado y variado considerando quién habla, a quién y para qué.',
            2: 'El vocabulario es limitado o repetitivo y pierde parcialmente de vista la situación o al destinatario.',
            1: 'El vocabulario es muy limitado o no corresponde con la situación comunicativa ni el destinatario.',
        },
        'puntuacion': {
            3: 'Usa correctamente la puntuación para construir el significado del texto.',
            2: 'Emplea algunos signos, pero omite otros o comete errores que afectan parcialmente la lectura.',
            1: 'No usa los signos elementales o los emplea de forma que altera el significado.',
        },
        'ortografia': {
            3: 'Aplica acentuación y reglas ortográficas, incluso en palabras con grafías que pueden confundirse.',
            2: 'Presenta algunos errores de acentuación y errores mínimos en palabras de grafía dudosa.',
            1: 'Muestra numerosos errores de acentuación y un descuido general de las reglas ortográficas.',
        },
    },
}


class SisatAssessment(core.db.Model):
    __tablename__ = 'sisat_assessment'
    id = core.db.Column(core.db.Integer, primary_key=True)
    student_id = core.db.Column(core.db.Integer, core.db.ForeignKey('student.id'), nullable=False, index=True)
    skill = core.db.Column(core.db.String(20), nullable=False, index=True)
    application_date = core.db.Column(core.db.Date, nullable=False, default=date.today, index=True)
    visit = core.db.Column(core.db.Integer, nullable=False, default=1)
    scores_json = core.db.Column(core.db.Text, default='{}')
    question_set_json = core.db.Column(core.db.Text, default='[]')
    total = core.db.Column(core.db.Integer, default=0)
    level = core.db.Column(core.db.String(30), default='REQUIERE APOYO')
    observations = core.db.Column(core.db.Text, default='')
    applied_by = core.db.Column(core.db.Integer, core.db.ForeignKey('user.id'), nullable=True)
    created_at = core.db.Column(core.db.DateTime, default=datetime.utcnow, nullable=False)
    updated_at = core.db.Column(core.db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow, nullable=False)
    __table_args__ = (
        core.db.UniqueConstraint('student_id', 'skill', 'application_date', 'visit', name='uq_sisat_application'),
    )


def _profile():
    try:
        import multi_user
        return multi_user._profile(session.get('uid'))
    except Exception:
        return None


def _can_capture():
    profile = _profile()
    return bool(profile and profile.active and profile.role in ('ADMIN', 'DIRECCION', 'DOCENTE'))


def _can_export():
    profile = _profile()
    return bool(profile and profile.active and profile.role in ('ADMIN', 'DIRECCION'))


def _level(skill, total):
    if skill in ('lectura', 'escritura'):
        return 'NIVEL ESPERADO' if total >= 15 else ('EN DESARROLLO' if total >= 10 else 'REQUIERE APOYO')
    return 'NIVEL ESPERADO' if total >= 8 else ('EN DESARROLLO' if total >= 5 else 'REQUIERE APOYO')


def _scores(row):
    try:
        value = json.loads(row.scores_json or '{}')
        return value if isinstance(value, dict) else {}
    except Exception:
        return {}


def _fraction(value):
    value = Fraction(value).limit_denominator()
    if value.denominator == 1:
        return str(value.numerator)
    return f'{value.numerator}/{value.denominator}'


def _mental_questions(grade, rng=None):
    rng = rng or random.SystemRandom()
    questions = []

    def add(prompt, answer, category):
        questions.append({'number': len(questions) + 1, 'prompt': prompt, 'answer': str(answer), 'category': category})

    if grade == '1':
        a = rng.randint(24, 69); b = rng.randint(16, 39); add(f'{a} más {b}', a + b, 'Suma de naturales')
        a = rng.randint(6, 14); b = rng.randint(4, 9); add(f'{a} por {b}', a * b, 'Multiplicación')
        a = rng.randint(12, 25); b = rng.randint(2, 4); c = rng.randint(8, 22); add(f'{a} por {b} menos {c}', a * b - c, 'Operaciones combinadas')
        den = rng.choice([4, 6, 8]); n1 = rng.randint(1, den - 2); n2 = rng.randint(1, den - n1); add(f'{n1}/{den} más {n2}/{den}', _fraction(Fraction(n1 + n2, den)), 'Suma de fracciones')
        den = rng.choice([2, 4, 5, 10]); num = rng.randint(1, den - 1); add(f'Convierte {num}/{den} en número decimal', f'{num/den:g}', 'Fracción a decimal')
        d1, d2 = rng.choice([(2, 3), (3, 4), (3, 5), (4, 5)]); n1 = rng.randint(1, d1 - 1); n2 = rng.randint(1, d2 - 1); add(f'{n1}/{d1} por {n2}/{d2}', _fraction(Fraction(n1*n2, d1*d2)), 'Multiplicación de fracciones')
        b = rng.choice([10.2, 11.5, 12.4, 15.25]); result = rng.randint(70, 130); a = result + b; add(f'{a:g} menos {b:g}', f'{result:g}', 'Resta de decimales')
        x = rng.randint(4, 11); d = rng.choice([0.25, 0.5, 0.75]); sub = rng.randint(2, 5); add(f'{x-d:g} más {d:g} menos {sub}', x-sub, 'Operaciones con decimales')
        a = rng.randint(4, 12); add(f'¿Cuánto es {a} al cuadrado?', a*a, 'Potencia cuadrada')
        den = rng.choice([2, 4]); frac = Fraction(rng.randint(1, den-1), den); whole = rng.randint(2, 5); sub = Fraction(rng.randint(1, 3), 3); add(f'{float(frac):g} más {whole} menos {_fraction(sub)}', _fraction(frac + whole - sub), 'Fracciones y decimales')
    elif grade == '2':
        h = rng.randint(1, 2); minutes = rng.choice([15, 25, 35, 45]); extra = rng.choice([20, 30, 40, 50]); total = h*60 + minutes + extra; add(f'¿Cuánto es {h} hora{'' if h == 1 else 's'} {minutes} minutos más {extra} minutos?', f'{total//60} h {total%60:02d} min', 'Suma de tiempo')
        divisor = rng.choice([3, 4, 5]); result = rng.randint(8, 20); add(f'{"Tercera" if divisor == 3 else "Cuarta" if divisor == 4 else "Quinta"} parte de {divisor*result}', result, 'Parte de una cantidad')
        pct = rng.choice([10, 20, 25, 50]); base = rng.choice([40, 60, 80, 90, 120, 160]); add(f'{pct}% de {base}', int(base*pct/100), 'Porcentaje')
        a = rng.randint(12, 35); b = rng.randint(2, a-3); add(f'{a}x menos {b}x', f'{a-b}x', 'Términos semejantes')
        start = rng.choice([0.2, 0.4, 0.6, 1.1]); step = rng.choice([0.2, 0.3, 0.5]); seq = [start + step*i for i in range(5)]; add(f'¿Qué números siguen en esta serie: {seq[0]:g}, {seq[1]:g}, {seq[2]:g}, __, __?', f'{seq[3]:g} y {seq[4]:g}', 'Sucesión decimal')
        result = rng.choice([20.5, 25.5, 30.5, 35.5]); b = rng.choice([11.5, 12.5, 13.5]); add(f'{result+b:g} menos {b:g}', f'{result:g}', 'Resta de decimales')
        den = rng.choice([2, 4, 5]); num = rng.randint(1, den-1); whole = den*rng.randint(8, 16); add(f'{num}/{den} de {whole}', int(num*whole/den), 'Fracción de una cantidad')
        decimal = rng.choice([0.25, 0.5, 0.75]); whole = rng.choice([40, 60, 80, 120]); add(f'{decimal:g} por {whole}', f'{decimal*whole:g}', 'Producto decimal')
        a = rng.randint(2, 5); add(f'¿Cuánto es {a} al cubo?', a**3, 'Potencia cúbica')
        n1, d1, n2, d2 = rng.choice([(1,2,1,4),(2,3,1,3),(3,4,1,2)]); add(f'{n1}/{d1} entre {n2}/{d2}', _fraction(Fraction(n1,d1)/Fraction(n2,d2)), 'División de fracciones')
    else:
        result = rng.choice([600, 700, 800, 900]); b = rng.randint(21, 79); add(f'{result-b} más {b}', result, 'Suma de naturales')
        a = rng.choice([600, 700, 800, 900]); b = rng.randint(41, 129); add(f'{a} menos {b}', a-b, 'Resta de naturales')
        a = rng.choice([20, 30, 40, 50, 60, 70]); b = rng.choice([200, 300, 400, 500]); add(f'{a} por {b}', a*b, 'Multiplicación')
        divisor = rng.randint(3, 9); first = divisor*rng.randint(5, 12); multiplier = rng.randint(3, 8); add(f'{first} entre {divisor} por {multiplier}', first//divisor*multiplier, 'División y multiplicación')
        a = rng.randint(3, 6); sub = rng.randint(2, 9); add(f'{a} al cubo, menos {sub}', a**3-sub, 'Potencia y resta')
        x = rng.randint(2, 12); coef = rng.randint(2, 7); add(f'¿Cuál es el valor de x en {coef}x menos {coef*x} = 0?', x, 'Ecuación lineal')
        values = rng.choice([(Fraction(1,2),Fraction(3,4),Fraction(2,8)),(Fraction(2,3),Fraction(1,2),Fraction(1,6)),(Fraction(3,4),Fraction(2,3),Fraction(5,12))]); add(f'{_fraction(values[0])} más {_fraction(values[1])} menos {_fraction(values[2])}', _fraction(values[0]+values[1]-values[2]), 'Operaciones con fracciones')
        frac = rng.choice([Fraction(1,2),Fraction(3,4),Fraction(2,5)]); dec = rng.choice([0.25,0.5,0.75]); add(f'{dec:g} más {_fraction(frac)}', _fraction(Fraction(str(dec))+frac), 'Decimal y fracción')
        den = rng.choice([3,5,7]); seq = [Fraction(2**i, den*(2**i)) for i in range(4)]; add(f'¿Qué fracciones siguen en esta serie: 1/{den}, 2/{den*2}, 4/{den*4}, __, __?', f'8/{den*8} y 16/{den*16}', 'Sucesión de fracciones equivalentes')
        a = rng.randint(30, 70); b = rng.randint(20, 100-a); add(f'Dos ángulos interiores de un triángulo miden {a}° y {b}°. ¿Cuánto mide el tercer ángulo?', f'{180-a-b}°', 'Ángulos de un triángulo')
    return questions


def _practice_questions(grade):
    return {
        '1': [('600 menos 500', '100'), ('¿Cuánto es la mitad de 62?', '31')],
        '2': [('15 por 10', '150'), ('¿60 entre qué número da 20?', '3')],
        '3': [('20 más 18', '38'), ('¿Qué número multiplicado por 5 da 40?', '8')],
    }[grade]


def _all_assessments():
    return SisatAssessment.query.execution_options(group_scope_disabled=True).all()


def _all_students():
    return core.Student.query.execution_options(group_scope_disabled=True).all()


def _student_groups():
    try:
        import group_workspaces
        rows = group_workspaces.RecordGroup.query.filter_by(entity_type='student').all()
        return {row.entity_id: row.group_code for row in rows}
    except Exception:
        return {}


def _active_group():
    try:
        import group_workspaces
        return group_workspaces.active_group_code() or '1A'
    except Exception:
        return '1A'


def _summary_rows(assessments, students):
    student_map = {student.id: student for student in students}
    groups = _student_groups()
    rows = []
    for assessment in assessments:
        student = student_map.get(assessment.student_id)
        if not student:
            continue
        rows.append({
            'group': groups.get(student.id, ''), 'list_no': student.list_no or '', 'student': student.full_name,
            'skill': SKILLS.get(assessment.skill, {}).get('title', assessment.skill),
            'date': assessment.application_date, 'visit': assessment.visit, 'scores': _scores(assessment),
            'total': assessment.total, 'level': assessment.level, 'observations': assessment.observations or '',
            'student_id': student.id,
        })
    return rows


def _parse_template_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ('%Y-%m-%d', '%d/%m/%Y'):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                continue
    raise ValueError('La fecha de aplicación no es válida.')


def _import_capture_workbook(file_storage):
    if not file_storage or not file_storage.filename:
        raise ValueError('Selecciona un archivo Excel.')
    if not file_storage.filename.lower().endswith('.xlsx'):
        raise ValueError('El archivo debe estar en formato .xlsx.')
    payload = file_storage.read(8 * 1024 * 1024 + 1)
    if len(payload) > 8 * 1024 * 1024:
        raise ValueError('El archivo excede el límite de 8 MB.')
    try:
        workbook = load_workbook(io.BytesIO(payload), data_only=False, read_only=False)
    except Exception as exc:
        raise ValueError('No fue posible abrir el archivo. Descarga un formato nuevo desde SiSAT.') from exc

    required = {'Lectura', 'Escritura', 'Cálculo mental'}
    if not required.issubset(workbook.sheetnames):
        raise ValueError('El archivo no corresponde al formato SiSAT generado por la plataforma.')
    students = core.Student.query.filter_by(status='ACTIVO').all()
    student_map = {student.id: student for student in students}
    imported = 0
    ignored = 0

    for sheet_name, skill_key in (('Lectura', 'lectura'), ('Escritura', 'escritura'), ('Cálculo mental', 'calculo')):
        ws = workbook[sheet_name]
        application_date = _parse_template_date(ws['H5'].value)
        try:
            visit = int(ws['K5'].value or 1)
        except (TypeError, ValueError) as exc:
            raise ValueError(f'La visita de la hoja {sheet_name} no es válida.') from exc
        if visit < 1 or visit > 9:
            raise ValueError(f'La visita de la hoja {sheet_name} debe estar entre 1 y 9.')
        for row_number in range(8, ws.max_row + 1):
            raw_id = ws.cell(row_number, 1).value
            if raw_id in (None, ''):
                continue
            try:
                student_id = int(raw_id)
            except (TypeError, ValueError):
                ignored += 1
                continue
            if student_id not in student_map:
                ignored += 1
                continue
            observations_col = 18 if skill_key != 'calculo' else 16
            observations = str(ws.cell(row_number, observations_col).value or '').strip()
            scores = {}
            if skill_key in ('lectura', 'escritura'):
                mapping = {'BUENA': 3, 'REGULAR': 2, 'INADECUADA': 1}
                values = [ws.cell(row_number, col).value for col in range(4, 15, 2)]
                if not any(value not in (None, '') for value in values) and not observations:
                    continue
                for (key, _, _), value in zip(SKILLS[skill_key]['components'], values):
                    normalized = str(value or '').strip().upper()
                    if normalized and normalized not in mapping:
                        raise ValueError(f'Valor no válido en {sheet_name}, fila {row_number}: {value}.')
                    scores[key] = mapping.get(normalized, 0)
                total = sum(scores.values())
            else:
                values = [ws.cell(row_number, col).value for col in range(4, 14)]
                if not any(value not in (None, '') for value in values) and not observations:
                    continue
                for number, value in enumerate(values, start=1):
                    normalized = str(value if value is not None else '').strip().upper()
                    if normalized not in ('', '0', '1', '1V'):
                        raise ValueError(f'Código no válido en Cálculo mental, fila {row_number}: {value}.')
                    scores[f'P{number}'] = normalized or '0'
                total = sum(1 for value in scores.values() if value in ('1', '1V'))
            assessment = SisatAssessment.query.execution_options(group_scope_disabled=True).filter_by(student_id=student_id, skill=skill_key, application_date=application_date, visit=visit).first()
            if not assessment:
                assessment = SisatAssessment(student_id=student_id, skill=skill_key, application_date=application_date, visit=visit)
            assessment.scores_json = json.dumps(scores, ensure_ascii=False)
            assessment.total = total
            assessment.level = _level(skill_key, total)
            assessment.observations = observations
            assessment.applied_by = session.get('uid')
            core.db.session.add(assessment)
            imported += 1
    if not imported:
        raise ValueError('El archivo no contiene resultados capturados.')
    core.db.session.commit()
    return imported, ignored


def _export_workbook(scope, skill='', student_id=None):
    assessments = _all_assessments()
    students = _all_students()
    groups = _student_groups()
    active_group = _active_group()
    if scope == 'group':
        allowed_ids = {student.id for student in students if groups.get(student.id) == active_group}
        assessments = [row for row in assessments if row.student_id in allowed_ids]
        students = [student for student in students if student.id in allowed_ids]
    elif scope == 'student':
        assessments = [row for row in assessments if row.student_id == student_id]
        students = [student for student in students if student.id == student_id]
    if skill in SKILLS:
        assessments = [row for row in assessments if row.skill == skill]
    rows = _summary_rows(assessments, students)

    output = io.BytesIO(); wb = xlsxwriter.Workbook(output, {'in_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False})
    school_config = core.cfg()
    wb.set_properties({'title': 'Informe de resultados SiSAT', 'subject': 'Gráficas y concentrados institucionales', 'author': school_config.school or 'Sistema Integral Escolar'})
    title = wb.add_format({'bold': True, 'font_size': 16, 'font_color': '#FFFFFF', 'bg_color': '#7B1024', 'align': 'center', 'valign': 'vcenter'})
    header = wb.add_format({'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#217346', 'align': 'center', 'valign': 'vcenter', 'text_wrap': True, 'border': 1})
    cell = wb.add_format({'border': 1, 'valign': 'top'}); center = wb.add_format({'border': 1, 'align': 'center', 'valign': 'top'})
    date_fmt = wb.add_format({'border': 1, 'align': 'center', 'num_format': 'dd/mm/yyyy'})
    percent_fmt = wb.add_format({'border': 1, 'align': 'center', 'num_format': '0.0%'})
    level_formats = {key: wb.add_format({'border': 1, 'align': 'center', 'bold': True, 'bg_color': color}) for key, color in LEVEL_COLORS.items()}

    ws = wb.add_worksheet('Resumen')
    ws.merge_range(0, 0, 0, 7, 'CONCENTRADO DE RESULTADOS SiSAT', title)
    ws.merge_range(1, 0, 1, 7, f'{school_config.school or ""} · CCT {school_config.cct or "Sin registro"} · Ciclo {school_config.cycle or ""} · Emitido {date.today().strftime("%d/%m/%Y")}', wb.add_format({'align': 'center', 'italic': True, 'font_color': '#657085'}))
    ws.write_row(2, 0, ['Ámbito', 'Alumnos evaluados', 'Aplicaciones', 'Nivel esperado', 'En desarrollo', 'Requiere apoyo', 'Promedio lectura/escritura', 'Promedio cálculo'], header)
    unique_students = len({row['student_id'] for row in rows}); levels = {level: sum(1 for row in rows if row['level'] == level) for level in LEVEL_COLORS}
    lit = [row['total'] for row in rows if row['skill'] != 'Cálculo mental']; mental = [row['total'] for row in rows if row['skill'] == 'Cálculo mental']
    label = 'Escuela' if scope == 'school' else (f'Grupo {active_group}' if scope == 'group' else (students[0].full_name if students else 'Alumno'))
    ws.write_row(3, 0, [label, unique_students, len(rows), levels['NIVEL ESPERADO'], levels['EN DESARROLLO'], levels['REQUIERE APOYO'], round(sum(lit)/len(lit), 2) if lit else '', round(sum(mental)/len(mental), 2) if mental else ''], center)
    ws.write(6, 0, 'Distribución por habilidad', header)
    ws.write_row(7, 0, ['Habilidad', 'Aplicaciones', 'Nivel esperado', 'En desarrollo', 'Requiere apoyo', 'Promedio'], header)
    out_row = 8
    for skill_key, config in SKILLS.items():
        subset = [row for row in rows if row['skill'] == config['title']]
        if not subset: continue
        ws.write_row(out_row, 0, [config['title'], len(subset), sum(1 for row in subset if row['level']=='NIVEL ESPERADO'), sum(1 for row in subset if row['level']=='EN DESARROLLO'), sum(1 for row in subset if row['level']=='REQUIERE APOYO'), round(sum(row['total'] for row in subset)/len(subset),2)], center); out_row += 1
    skill_start_row = 8
    skill_end_row = out_row - 1
    if scope == 'school':
        out_row += 2; ws.write(out_row, 0, 'Resultados por grupo', header); out_row += 1
        ws.write_row(out_row, 0, ['Grupo', 'Alumnos evaluados', 'Aplicaciones', 'Nivel esperado', 'En desarrollo', 'Requiere apoyo', 'Promedio'], header); out_row += 1
        for group_code in sorted({row['group'] for row in rows if row['group']}):
            subset = [row for row in rows if row['group'] == group_code]
            ws.write_row(out_row, 0, [group_code, len({row['student_id'] for row in subset}), len(subset), sum(1 for row in subset if row['level']=='NIVEL ESPERADO'), sum(1 for row in subset if row['level']=='EN DESARROLLO'), sum(1 for row in subset if row['level']=='REQUIERE APOYO'), round(sum(row['total'] for row in subset)/len(subset),2)], center); out_row += 1
    if rows:
        distribution = wb.add_chart({'type': 'doughnut'})
        distribution.add_series({
            'name': 'Distribución general por nivel',
            'categories': '=Resumen!$D$3:$F$3',
            'values': '=Resumen!$D$4:$F$4',
            'points': [{'fill': {'color': '#63BE7B'}}, {'fill': {'color': '#FFCE54'}}, {'fill': {'color': '#F8696B'}}],
            'data_labels': {'percentage': True, 'category': True},
        })
        distribution.set_title({'name': 'Distribución general por nivel'})
        distribution.set_legend({'none': True})
        distribution.set_hole_size(48)
        distribution.set_style(10)
        ws.insert_chart('J2', distribution, {'x_scale': 1.12, 'y_scale': 1.05})
    if skill_end_row >= skill_start_row:
        skills_chart = wb.add_chart({'type': 'column'})
        for col, name, color in ((2, 'Nivel esperado', '#63BE7B'), (3, 'En desarrollo', '#FFCE54'), (4, 'Requiere apoyo', '#F8696B')):
            skills_chart.add_series({
                'name': name,
                'categories': ['Resumen', skill_start_row, 0, skill_end_row, 0],
                'values': ['Resumen', skill_start_row, col, skill_end_row, col],
                'fill': {'color': color}, 'border': {'none': True},
            })
        skills_chart.set_title({'name': 'Resultados por habilidad'})
        skills_chart.set_y_axis({'name': 'Aplicaciones', 'major_gridlines': {'visible': False}, 'min': 0})
        skills_chart.set_legend({'position': 'top'})
        skills_chart.set_style(10)
        ws.insert_chart('J19', skills_chart, {'x_scale': 1.22, 'y_scale': 1.1})
    ws.set_column(0, 0, 30); ws.set_column(1, 7, 18); ws.set_column(9, 16, 13); ws.freeze_panes(3, 0)
    ws.set_landscape(); ws.fit_to_pages(1, 1); ws.set_header(f'&C&BSiSAT · {school_config.school or ""}'); ws.set_footer('&LConfidencial&C&P de &N&R&D')

    detail = wb.add_worksheet('Detalle')
    detail.merge_range(0, 0, 0, 9, 'DETALLE DE RESULTADOS SiSAT', title)
    headers = ['Grupo', 'No.', 'Alumno', 'Habilidad', 'Fecha', 'Visita', 'Puntaje', 'Resultado', 'Resultados por componente/pregunta', 'Observaciones']
    detail.write_row(2, 0, headers, header)
    for index, row in enumerate(rows, start=3):
        score_text = ', '.join(f'{key}: {value}' for key, value in row['scores'].items())
        values = [row['group'], row['list_no'], row['student'], row['skill'], row['date'], row['visit'], row['total'], row['level'], score_text, row['observations']]
        for col, value in enumerate(values):
            fmt = date_fmt if col == 4 else (level_formats.get(row['level'], center) if col == 7 else (center if col in (0,1,5,6) else cell))
            detail.write(index, col, value, fmt)
    detail.autofilter(2, 0, max(3, len(rows)+2), 9); detail.freeze_panes(3, 3)
    detail.set_column(0, 1, 10); detail.set_column(2, 3, 28); detail.set_column(4, 7, 16); detail.set_column(8, 9, 48)
    detail.conditional_format(3, 6, max(3, len(rows)+2), 6, {'type': '3_color_scale', 'min_color': '#FADDDD', 'mid_color': '#FFF0C9', 'max_color': '#DFF2E5'})

    analysis = wb.add_worksheet('Componentes')
    analysis.merge_range(0, 0, 0, 6, 'ANÁLISIS POR COMPONENTE Y PREGUNTA', title)
    analysis.write_row(2, 0, ['Habilidad', 'Componente / pregunta', 'Registros', 'Puntaje obtenido', 'Puntaje posible', 'Porcentaje', 'Interpretación'], header)
    analysis_row = 3
    for skill_key, config in SKILLS.items():
        subset = [row for row in rows if row['skill'] == config['title']]
        if not subset: continue
        if skill_key in ('lectura', 'escritura'):
            measures = [(key, label, 3) for key, label, _ in config['components']]
        else:
            measures = [(f'P{number}', f'Pregunta {number}', 1) for number in range(1, 11)]
        for key, label_text, maximum in measures:
            values = [row['scores'].get(key) for row in subset if row['scores'].get(key) not in (None, '', 0, '0')]
            if skill_key == 'calculo':
                obtained = sum(1 for value in values if value in ('1', '1V')); count = len(subset); possible = count
                visual = sum(1 for value in values if value == '1V')
            else:
                numeric = [int(value) for value in values if str(value).isdigit()]; obtained = sum(numeric); count = len(numeric); possible = count * maximum; visual = 0
            percentage = obtained / possible if possible else 0
            if skill_key == 'calculo':
                interpretation = 'Desarrollo generalizado' if percentage >= .81 else ('Diferencias de desempeño' if percentage >= .51 else 'Dificultad generalizada')
                if percentage >= .81 and obtained and visual / obtained >= .5: interpretation += '; reforzar presentación verbal'
            else:
                interpretation = 'Manejo adecuado' if percentage >= .81 else ('Avance significativo' if percentage >= .56 else 'Poco avance')
            values_out = [config['title'], label_text, count, obtained, possible, percentage, interpretation]
            for col, value in enumerate(values_out): analysis.write(analysis_row, col, value, percent_fmt if col == 5 else (center if col in (2,3,4) else cell))
            analysis.set_row(analysis_row, 28); analysis_row += 1
    analysis.set_column(0, 1, 30); analysis.set_column(2, 4, 16); analysis.set_column(5, 5, 14); analysis.set_column(6, 6, 38)
    analysis.autofilter(2, 0, max(3, analysis_row-1), 6); analysis.freeze_panes(3, 2)

    rubric_sheet = wb.add_worksheet('Rúbricas')
    rubric_sheet.merge_range(0, 0, 0, 5, 'RÚBRICAS DE LECTURA Y PRODUCCIÓN DE TEXTOS', title)
    rubric_sheet.write_row(2, 0, ['Habilidad', 'Componente', '3 puntos', '2 puntos', '1 punto', 'Rango global'], header)
    rubric_row = 3
    for skill_key in ('lectura', 'escritura'):
        config = SKILLS[skill_key]
        for key, label_text, descriptors in config['components']:
            rubric_sheet.write_row(rubric_row, 0, [config['title'], label_text, f'{descriptors[0]}. {RUBRIC_GUIDANCE[skill_key][key][3]}', f'{descriptors[1]}. {RUBRIC_GUIDANCE[skill_key][key][2]}', f'{descriptors[2]}. {RUBRIC_GUIDANCE[skill_key][key][1]}', '15–18: Nivel esperado · 10–14: En desarrollo · 9 o menos: Requiere apoyo'], cell)
            rubric_sheet.set_row(rubric_row, 72); rubric_row += 1
    rubric_sheet.write_row(rubric_row, 0, ['Cálculo mental', 'Códigos y nivel', '1 = correcta sin apoyo visual', '1V = correcta con apoyo visual', '0 = equivocada o sin respuesta', '8–10: Nivel esperado · 5–7: En desarrollo · 0–4: Requiere apoyo'], cell)
    rubric_sheet.set_row(rubric_row, 54)
    rubric_sheet.set_column(0, 1, 28); rubric_sheet.set_column(2, 4, 46); rubric_sheet.set_column(5, 5, 34); rubric_sheet.freeze_panes(3, 2)
    wb.close(); output.seek(0); return output


def _capture_template_workbook(students, group_code):
    """Build an Excel capture template with formulas and official SiSAT bands."""
    output = io.BytesIO()
    wb = xlsxwriter.Workbook(output, {'in_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False})
    wb.set_properties({
        'title': f'Formato automatizado SiSAT · Grupo {group_code}',
        'subject': 'Captura de resultados de lectura, escritura y cálculo mental',
        'author': 'Sistema Integral Escolar',
        'comments': 'Las celdas amarillas son editables; puntajes y niveles se calculan automáticamente.',
    })
    config = core.cfg()
    last_data_row = max(7, 7 + len(students) - 1)

    title = wb.add_format({'bold': True, 'font_size': 15, 'font_color': '#172033', 'bottom': 2, 'bottom_color': '#7B1024'})
    subtitle = wb.add_format({'italic': True, 'font_color': '#657085'})
    meta_label = wb.add_format({'bold': True, 'font_color': '#7B1024'})
    meta_value = wb.add_format({'bottom': 1, 'bottom_color': '#AAB4C3'})
    input_meta = wb.add_format({'bg_color': '#FFF0C9', 'border': 1, 'border_color': '#D8B861', 'locked': False, 'num_format': 'dd/mm/yyyy'})
    input_visit = wb.add_format({'bg_color': '#FFF0C9', 'border': 1, 'border_color': '#D8B861', 'locked': False, 'align': 'center'})
    header = wb.add_format({'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#217346', 'align': 'center', 'valign': 'vcenter', 'text_wrap': True, 'border': 1, 'border_color': '#FFFFFF'})
    subheader = wb.add_format({'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#4F7D65', 'align': 'center', 'valign': 'vcenter', 'text_wrap': True, 'border': 1, 'border_color': '#FFFFFF'})
    identity = wb.add_format({'valign': 'vcenter', 'bottom': 1, 'bottom_color': '#DDE3EA'})
    input_text = wb.add_format({'bg_color': '#FFF8DC', 'align': 'center', 'valign': 'vcenter', 'locked': False, 'bottom': 1, 'bottom_color': '#E8D89A'})
    input_notes = wb.add_format({'bg_color': '#FFF8DC', 'valign': 'top', 'text_wrap': True, 'locked': False, 'bottom': 1, 'bottom_color': '#E8D89A'})
    formula_num = wb.add_format({'align': 'center', 'valign': 'vcenter', 'bold': True, 'bg_color': '#F1F5F9', 'bottom': 1, 'bottom_color': '#DDE3EA'})
    formula_level = wb.add_format({'align': 'center', 'valign': 'vcenter', 'bold': True, 'text_wrap': True, 'bg_color': '#F1F5F9', 'bottom': 1, 'bottom_color': '#DDE3EA'})
    note = wb.add_format({'font_color': '#657085', 'italic': True, 'text_wrap': True})
    rubric_cell = wb.add_format({'valign': 'top', 'text_wrap': True, 'bottom': 1, 'bottom_color': '#DDE3EA'})
    pct = wb.add_format({'align': 'center', 'num_format': '0.0%'})
    integer = wb.add_format({'align': 'center', 'num_format': '0'})

    def setup_sheet(ws, skill_title, last_col):
        ws.hide_gridlines(2)
        ws.merge_range(0, 1, 0, last_col, f'FORMATO AUTOMATIZADO SiSAT · {skill_title.upper()}', title)
        ws.merge_range(1, 1, 1, last_col, 'Capture únicamente las celdas amarillas. Los puntajes y el nivel se calculan de forma automática.', subtitle)
        ws.write(3, 1, 'Escuela:', meta_label); ws.merge_range(3, 2, 3, 5, config.school or '', meta_value)
        ws.write(3, 6, 'CCT:', meta_label); ws.merge_range(3, 7, 3, 8, config.cct or '', meta_value)
        ws.write(4, 1, 'Ciclo:', meta_label); ws.write(4, 2, config.cycle or '', meta_value)
        ws.write(4, 3, 'Grupo:', meta_label); ws.write(4, 4, group_code, meta_value)
        ws.write(4, 6, 'Fecha:', meta_label); ws.write_datetime(4, 7, datetime.combine(date.today(), datetime.min.time()), input_meta)
        ws.write(4, 9, 'Visita:', meta_label); ws.write_number(4, 10, 1, input_visit)
        ws.data_validation(4, 10, 4, 10, {'validate': 'integer', 'criteria': 'between', 'minimum': 1, 'maximum': 9, 'input_title': 'Número de visita', 'input_message': 'Capture un valor del 1 al 9.'})
        ws.set_row(0, 25); ws.set_row(1, 30); ws.set_row(6, 48)
        ws.freeze_panes(7, 3)
        ws.autofilter(6, 1, last_data_row, last_col)
        ws.set_landscape(); ws.fit_to_pages(1, 0); ws.repeat_rows(0, 6)

    def add_level_formatting(ws, col):
        for level, color in LEVEL_COLORS.items():
            ws.conditional_format(7, col, last_data_row, col, {'type': 'text', 'criteria': 'containing', 'value': level, 'format': wb.add_format({'bg_color': color, 'bold': True, 'font_color': '#172033'})})

    for skill_key, sheet_name in (('lectura', 'Lectura'), ('escritura', 'Escritura')):
        ws = wb.add_worksheet(sheet_name)
        setup_sheet(ws, SKILLS[skill_key]['title'], 17)
        headers = ['ID', 'No.', 'Alumno']
        for _, label, _ in SKILLS[skill_key]['components']:
            headers.extend([label, 'Puntos'])
        headers.extend(['Total', 'Resultado', 'Observaciones'])
        ws.write_row(6, 0, headers, header)
        for component_index in range(6):
            ws.write(6, 3 + component_index * 2, headers[3 + component_index * 2], header)
            ws.write(6, 4 + component_index * 2, 'Puntos', subheader)
        for offset, student in enumerate(students):
            row = 7 + offset; excel_row = row + 1
            ws.write_number(row, 0, student.id, identity)
            ws.write(row, 1, student.list_no or '', identity)
            ws.write(row, 2, student.full_name, identity)
            score_cells = []
            for index in range(6):
                input_col = 3 + index * 2; score_col = input_col + 1
                input_letter = xlsxwriter.utility.xl_col_to_name(input_col)
                score_letter = xlsxwriter.utility.xl_col_to_name(score_col)
                ws.write_blank(row, input_col, None, input_text)
                ws.write_formula(row, score_col, f'=IF({input_letter}{excel_row}="","",IF({input_letter}{excel_row}="BUENA",3,IF({input_letter}{excel_row}="REGULAR",2,1)))', formula_num, '')
                score_cells.append(f'{score_letter}{excel_row}')
            input_refs = ','.join(f'{xlsxwriter.utility.xl_col_to_name(3 + i * 2)}{excel_row}' for i in range(6))
            ws.write_formula(row, 15, f'=IF(COUNTA({input_refs})=0,"",SUM({",".join(score_cells)}))', formula_num, '')
            ws.write_formula(row, 16, f'=IF(P{excel_row}="","",IF(P{excel_row}>=15,"NIVEL ESPERADO",IF(P{excel_row}>=10,"EN DESARROLLO","REQUIERE APOYO")))', formula_level, '')
            ws.write_blank(row, 17, None, input_notes)
            ws.set_row(row, 34)
        # Apply validation only to the six editable rubric columns.
        for col in range(3, 14, 2):
            ws.data_validation(7, col, last_data_row, col, {'validate': 'list', 'source': ['BUENA', 'REGULAR', 'INADECUADA'], 'input_title': 'Valoración', 'input_message': 'Elija BUENA, REGULAR o INADECUADA.', 'error_title': 'Valor no permitido', 'error_message': 'Seleccione una opción de la lista.'})
        ws.set_column(0, 0, 2, None, {'hidden': True}); ws.set_column(1, 1, 7); ws.set_column(2, 2, 32)
        for col in range(3, 15, 2): ws.set_column(col, col, 19)
        for col in range(4, 15, 2): ws.set_column(col, col, 8)
        ws.set_column(15, 15, 8); ws.set_column(16, 16, 19); ws.set_column(17, 17, 32)
        add_level_formatting(ws, 16)
        ws.protect('', {'select_locked_cells': True, 'select_unlocked_cells': True, 'sort': True, 'autofilter': True})

    mental = wb.add_worksheet('Cálculo mental')
    setup_sheet(mental, SKILLS['calculo']['title'], 15)
    mental.write_row(6, 0, ['ID', 'No.', 'Alumno'] + [f'P{i}' for i in range(1, 11)] + ['Total', 'Resultado', 'Observaciones'], header)
    for offset, student in enumerate(students):
        row = 7 + offset; excel_row = row + 1
        mental.write_number(row, 0, student.id, identity); mental.write(row, 1, student.list_no or '', identity); mental.write(row, 2, student.full_name, identity)
        for col in range(3, 13): mental.write_blank(row, col, None, input_text)
        mental.write_formula(row, 13, f'=IF(COUNTA(D{excel_row}:M{excel_row})=0,"",COUNTIF(D{excel_row}:M{excel_row},"1")+COUNTIF(D{excel_row}:M{excel_row},"1V"))', formula_num, '')
        mental.write_formula(row, 14, f'=IF(N{excel_row}="","",IF(N{excel_row}>=8,"NIVEL ESPERADO",IF(N{excel_row}>=5,"EN DESARROLLO","REQUIERE APOYO")))', formula_level, '')
        mental.write_blank(row, 15, None, input_notes); mental.set_row(row, 34)
    mental.data_validation(7, 3, last_data_row, 12, {'validate': 'list', 'source': ['1', '1V', '0'], 'input_title': 'Código', 'input_message': '1 = correcta; 1V = correcta con apoyo visual; 0 = incorrecta.'})
    mental.set_column(0, 0, 2, None, {'hidden': True}); mental.set_column(1, 1, 7); mental.set_column(2, 2, 32); mental.set_column(3, 12, 7); mental.set_column(13, 13, 8); mental.set_column(14, 14, 19); mental.set_column(15, 15, 32)
    add_level_formatting(mental, 14)
    mental.protect('', {'select_locked_cells': True, 'select_unlocked_cells': True, 'sort': True, 'autofilter': True})

    summary = wb.add_worksheet('Resumen')
    summary.hide_gridlines(2); summary.set_tab_color('#7B1024')
    summary.merge_range('A1:G1', 'RESUMEN AUTOMÁTICO DE RESULTADOS SiSAT', title)
    summary.merge_range('A2:G2', f'{config.school or ""} · {config.cct or "Sin CCT"} · Ciclo {config.cycle or ""} · Grupo {group_code}', subtitle)
    summary.write_row('A5', ['Habilidad', 'Alumnos capturados', 'Nivel esperado', 'En desarrollo', 'Requiere apoyo', '% esperado', 'Promedio'], header)
    summary_rows = [('Lectura', 'Q', 'P'), ('Escritura', 'Q', 'P'), ('Cálculo mental', 'O', 'N')]
    for row_index, (sheet_name, level_col, total_col) in enumerate(summary_rows, start=5):
        first_excel, last_excel = 8, last_data_row + 1
        summary.write(row_index, 0, sheet_name, identity)
        summary.write_formula(row_index, 1, f'=COUNT({sheet_name!r}!{total_col}{first_excel}:{total_col}{last_excel})', integer, 0)
        summary.write_formula(row_index, 2, f'=COUNTIF({sheet_name!r}!{level_col}{first_excel}:{level_col}{last_excel},"NIVEL ESPERADO")', integer, 0)
        summary.write_formula(row_index, 3, f'=COUNTIF({sheet_name!r}!{level_col}{first_excel}:{level_col}{last_excel},"EN DESARROLLO")', integer, 0)
        summary.write_formula(row_index, 4, f'=COUNTIF({sheet_name!r}!{level_col}{first_excel}:{level_col}{last_excel},"REQUIERE APOYO")', integer, 0)
        summary.write_formula(row_index, 5, f'=IF(B{row_index+1}=0,"",C{row_index+1}/B{row_index+1})', pct, '')
        summary.write_formula(row_index, 6, f'=IF(B{row_index+1}=0,"",AVERAGE({sheet_name!r}!{total_col}{first_excel}:{total_col}{last_excel}))', wb.add_format({'align': 'center', 'num_format': '0.00'}), '')
    summary.write('A11', 'Lectura y escritura: 15–18 Nivel esperado · 10–14 En desarrollo · 0–9 Requiere apoyo.', note)
    summary.write('A12', 'Cálculo mental: 8–10 Nivel esperado · 5–7 En desarrollo · 0–4 Requiere apoyo.', note)
    summary.set_column('A:A', 28); summary.set_column('B:E', 18); summary.set_column('F:G', 14); summary.set_row(4, 34)
    summary.conditional_format('C6:C8', {'type': 'data_bar', 'bar_color': '#63BE7B'}); summary.conditional_format('E6:E8', {'type': 'data_bar', 'bar_color': '#F8696B'})

    rubric = wb.add_worksheet('Rúbricas')
    rubric.hide_gridlines(2); rubric.set_tab_color('#657085')
    rubric.merge_range('A1:F1', 'RÚBRICAS Y CRITERIOS DE VALORACIÓN', title)
    rubric.write_row('A3', ['Habilidad', 'Componente', 'BUENA · 3 puntos', 'REGULAR · 2 puntos', 'INADECUADA · 1 punto', 'Rango global'], header)
    rubric_row = 3
    for skill_key in ('lectura', 'escritura'):
        for key, label_text, descriptors in SKILLS[skill_key]['components']:
            rubric.write_row(rubric_row, 0, [SKILLS[skill_key]['title'], label_text, f'{descriptors[0]}. {RUBRIC_GUIDANCE[skill_key][key][3]}', f'{descriptors[1]}. {RUBRIC_GUIDANCE[skill_key][key][2]}', f'{descriptors[2]}. {RUBRIC_GUIDANCE[skill_key][key][1]}', '15–18: Nivel esperado\n10–14: En desarrollo\n0–9: Requiere apoyo'], rubric_cell)
            rubric.set_row(rubric_row, 88); rubric_row += 1
    rubric.write_row(rubric_row, 0, ['Cálculo mental', 'Códigos', '1 = correcta sin apoyo visual', '1V = correcta con apoyo visual', '0 = equivocada o sin respuesta', '8–10: Nivel esperado\n5–7: En desarrollo\n0–4: Requiere apoyo'], rubric_cell)
    rubric.set_row(rubric_row, 58); rubric.set_column('A:B', 26); rubric.set_column('C:E', 47); rubric.set_column('F:F', 28); rubric.freeze_panes(3, 2)

    # Open the workbook on its reader-facing summary.
    summary.activate(); summary.select()
    wb.close(); output.seek(0)
    return output


def install(app):
    try:
        with app.app_context(): core.db.create_all()
    except Exception:
        pass

    @app.route('/sisat')
    def sisat_dashboard():
        if not session.get('uid'): return redirect('/login')
        group = _active_group(); students = core.Student.query.filter_by(status='ACTIVO').order_by(core.Student.list_no, core.Student.paternal).all()
        student_ids = [student.id for student in students]
        rows = SisatAssessment.query.filter(SisatAssessment.student_id.in_(student_ids)).all() if student_ids else []
        cards = ''
        for key, config in SKILLS.items():
            subset = [row for row in rows if row.skill == key]
            completed = len({row.student_id for row in subset}); priority = sum(1 for row in subset if row.level == 'REQUIERE APOYO')
            cards += f'''<div class="card"><h2>{escape(config['title'])}</h2><div class="kpi">{completed}/{len(students)}</div><p class="muted">Alumnos con registro · {priority} aplicaciones requieren apoyo</p><a href="/sisat/capture/{key}" class="sisat-btn">Capturar resultados</a></div>'''
        export_block = ''
        if _can_export():
            student_options = ''.join(f'<option value="{s.id}">{escape(s.full_name)}</option>' for s in students)
            export_block = f'''<div class="card"><h2>Informe gráfico y concentrado</h2><p class="muted">Administración y Dirección pueden generar un archivo institucional con encabezado predeterminado, gráficas y resultados por escuela, grupo o alumno.</p><form method="get" action="/sisat/export.xlsx" class="grid"><label>Ámbito<select name="scope" id="sisat-scope" onchange="document.getElementById('sisat-student').style.display=this.value==='student'?'block':'none'"><option value="school">Toda la escuela</option><option value="group">Grupo {group}</option><option value="student">Alumno</option></select></label><label>Habilidad<select name="skill"><option value="">Todas</option>{''.join(f'<option value="{k}">{escape(v["title"])}</option>' for k,v in SKILLS.items())}</select></label><label id="sisat-student" style="display:none">Alumno<select name="student_id">{student_options}</select></label><div><button>Generar informe Excel</button></div></form></div>'''
        template_block = ''
        if _can_capture():
            template_block = f'''<div class="card"><h2>Formato automatizado de captura</h2><p>Descarga un Excel del grupo {group} con los alumnos precargados, listas desplegables, puntajes, niveles y resumen automático.</p><a href="/sisat/formato-captura.xlsx" class="sisat-btn">Descargar formato Excel</a><p class="muted">Incluye lectura, escritura, cálculo mental y las rúbricas oficiales.</p><hr style="border:0;border-top:1px solid #e8edf4;margin:18px 0"><h3>Importar resultados</h3><p class="muted">Después de llenar el formato, selecciónalo para guardar los resultados en la plataforma.</p><form method="post" action="/sisat/importar-formato" enctype="multipart/form-data" class="grid"><label>Archivo Excel<input type="file" name="file" accept=".xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" required></label><div><button>Importar resultados</button></div></form></div>'''
        materials_block = ''
        if _can_capture():
            materials_block = '''<div class="card"><h2>Actividades nuevas en PDF</h2><p>Genera lecturas originales con preguntas graduadas o ejercicios de escritura que permiten observar todos los componentes de las rúbricas.</p><a href="/sisat/materiales" class="sisat-btn">Generar actividades PDF</a></div>'''
        body = f'''<h1>SiSAT · Grupo {group}</h1><p class="muted">Exploración de habilidades básicas en lectura, producción de textos escritos y cálculo mental.</p><div class="grid">{cards}</div>{template_block}{materials_block}<div class="card"><h2>Prueba de cálculo mental</h2><p>Genera diez reactivos aleatorios con el mismo nivel y tipo de habilidad que los instrumentos oficiales de cada grado.</p><a href="/sisat/mental-test" class="sisat-btn">Generar prueba</a></div>{export_block}<style>.sisat-btn{{display:inline-block;background:#7b1024;color:white;text-decoration:none;padding:10px 14px;border-radius:9px;font-weight:800}}</style>'''
        return core.page('SiSAT', body)

    @app.route('/sisat/formato-captura.xlsx')
    def sisat_capture_template():
        if not session.get('uid'): return redirect('/login')
        if not _can_capture():
            flash('Tu rol no permite generar formatos de captura SiSAT.')
            return redirect('/sisat')
        group = _active_group()
        students = core.Student.query.filter_by(status='ACTIVO').order_by(core.Student.list_no, core.Student.paternal).all()
        output = _capture_template_workbook(students, group)
        safe_group = ''.join(char for char in group if char.isalnum() or char in ('-', '_')) or 'grupo'
        return send_file(output, as_attachment=True, download_name=f'formato_sisat_{safe_group}_{date.today().isoformat()}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.route('/sisat/importar-formato', methods=['POST'])
    def sisat_import_template():
        if not session.get('uid'): return redirect('/login')
        if not _can_capture():
            flash('Tu rol no permite importar resultados SiSAT.')
            return redirect('/sisat')
        try:
            imported, ignored = _import_capture_workbook(request.files.get('file'))
        except ValueError as exc:
            core.db.session.rollback()
            flash(str(exc))
            return redirect('/sisat')
        except Exception:
            core.db.session.rollback()
            flash('No fue posible importar el archivo. Verifica que sea un formato SiSAT válido.')
            return redirect('/sisat')
        message = f'Se importaron {imported} registros SiSAT correctamente.'
        if ignored:
            message += f' Se omitieron {ignored} filas que no pertenecen al grupo activo.'
        flash(message)
        return redirect('/sisat')

    @app.route('/sisat/capture/<skill>', methods=['GET', 'POST'])
    def sisat_capture(skill):
        if not session.get('uid'): return redirect('/login')
        if skill not in SKILLS: return redirect('/sisat')
        if not _can_capture(): flash('Tu rol no permite capturar resultados SiSAT.'); return redirect('/sisat')
        students = core.Student.query.filter_by(status='ACTIVO').order_by(core.Student.list_no, core.Student.paternal).all()
        app_date = request.values.get('date') or str(date.today()); visit = request.values.get('visit', type=int) or 1
        try: parsed_date = datetime.strptime(app_date, '%Y-%m-%d').date()
        except ValueError: parsed_date = date.today(); app_date = str(parsed_date)
        existing = {row.student_id: row for row in SisatAssessment.query.filter_by(skill=skill, application_date=parsed_date, visit=visit).all()}
        if request.method == 'POST':
            for student in students:
                if request.form.get(f'enabled_{student.id}') != '1': continue
                scores = {}
                if skill in ('lectura', 'escritura'):
                    for key, _, _ in SKILLS[skill]['components']:
                        value = request.form.get(f'{key}_{student.id}', type=int)
                        scores[key] = value if value in (1,2,3) else 0
                    total = sum(scores.values())
                else:
                    for number in range(1, 11):
                        code = request.form.get(f'q{number}_{student.id}', '0')
                        scores[f'P{number}'] = code if code in ('1', '1V', '0') else '0'
                    total = sum(1 for value in scores.values() if value in ('1','1V'))
                row = existing.get(student.id) or SisatAssessment(student_id=student.id, skill=skill, application_date=parsed_date, visit=visit)
                row.scores_json = json.dumps(scores, ensure_ascii=False); row.total = total; row.level = _level(skill, total)
                row.observations = request.form.get(f'obs_{student.id}', '').strip(); row.applied_by = session.get('uid')
                core.db.session.add(row)
            core.db.session.commit(); flash(f'Resultados de {SKILLS[skill]["title"]} guardados.'); return redirect(f'/sisat/capture/{skill}?date={app_date}&visit={visit}')
        headers = ''; rows_html = ''
        if skill in ('lectura', 'escritura'):
            headers = ''.join(f'<th>{escape(label)}</th>' for _, label, _ in SKILLS[skill]['components'])
        else: headers = ''.join(f'<th>P{number}</th>' for number in range(1,11))
        for student in students:
            row = existing.get(student.id); values = _scores(row) if row else {}; cells = ''
            if skill in ('lectura', 'escritura'):
                for key, _, descriptors in SKILLS[skill]['components']:
                    selected = int(values.get(key, 0) or 0)
                    options = '<option value="0">Sin valorar</option>' + ''.join(f'<option value="{score}" {"selected" if selected==score else ""}>{score} · {escape(descriptors[3-score])}</option>' for score in (3,2,1))
                    cells += f'<td><select name="{key}_{student.id}" style="min-width:190px">{options}</select></td>'
            else:
                for number in range(1,11):
                    selected = values.get(f'P{number}', '0'); options = ''.join(f'<option value="{code}" {"selected" if selected==code else ""}>{label}</option>' for code,label in [('1','1 · Correcta sin apoyo'),('1V','1V · Correcta con apoyo visual'),('0','0 · Incorrecta o sin respuesta')]); cells += f'<td><select name="q{number}_{student.id}" style="min-width:130px">{options}</select></td>'
            total = row.total if row else 0; level = row.level if row else 'SIN REGISTRO'; obs = row.observations if row else ''
            rows_html += f'''<tr><td><input class="sisat-enable" type="checkbox" name="enabled_{student.id}" value="1" {"checked" if row else ""} style="width:auto"></td><td>{student.list_no or ''}</td><td class="sticky-student"><b>{escape(student.full_name)}</b></td>{cells}<td><b>{total}</b></td><td>{escape(level)}</td><td><textarea name="obs_{student.id}" rows="2" style="min-width:210px">{escape(obs)}</textarea></td></tr>'''
        rubric = ''
        if skill in ('lectura','escritura'):
            rubric = '<details class="card"><summary><b>Consultar rúbrica de valoración</b></summary><div class="grid" style="margin-top:14px">' + ''.join(f'<div><h3>{escape(label)}</h3><p><b>3 · {escape(desc[0])}:</b> {escape(RUBRIC_GUIDANCE[skill][key][3])}<br><br><b>2 · {escape(desc[1])}:</b> {escape(RUBRIC_GUIDANCE[skill][key][2])}<br><br><b>1 · {escape(desc[2])}:</b> {escape(RUBRIC_GUIDANCE[skill][key][1])}</p></div>' for key,label,desc in SKILLS[skill]['components']) + '</div></details>'
        else:
            rubric = '<div class="card"><b>Códigos:</b> 1 = respuesta correcta sin presentación visual · 1V = respuesta correcta con presentación visual · 0 = respuesta equivocada o sin respuesta. Detenga la aplicación después de seis errores consecutivos.</div>'
        body = f'''<h1>{escape(SKILLS[skill]['title'])}</h1><form method="get" class="card grid"><label>Fecha<input type="date" name="date" value="{app_date}"></label><label>Visita<input type="number" name="visit" min="1" max="9" value="{visit}"></label><div><button>Cargar aplicación</button></div></form>{rubric}<form method="post" class="card"><input type="hidden" name="date" value="{app_date}"><input type="hidden" name="visit" value="{visit}"><div style="display:flex;gap:10px;margin-bottom:12px"><button type="button" style="width:auto" onclick="document.querySelectorAll('.sisat-enable').forEach(x=>x.checked=true)">Seleccionar todos</button><button type="button" style="width:auto;background:#6d737c" onclick="document.querySelectorAll('.sisat-enable').forEach(x=>x.checked=false)">Quitar selección</button></div><p class="muted">Solo se guardarán los alumnos seleccionados en la primera columna.</p><div class="scroll"><table class="sisat-capture"><tr><th>Aplicar</th><th>No.</th><th>Alumno</th>{headers}<th>Total</th><th>Resultado</th><th>Observaciones</th></tr>{rows_html}</table></div><br><button>Guardar resultados</button></form><style>.sisat-capture{{min-width:{'1800' if skill != 'calculo' else '2300'}px}}.sticky-student{{position:sticky;left:0;background:white;z-index:2;min-width:220px}}</style>'''
        return core.page(SKILLS[skill]['title'], body)

    @app.route('/sisat/mental-test')
    def sisat_mental_test():
        if not session.get('uid'): return redirect('/login')
        grade = request.args.get('grade', '1'); grade = grade if grade in ('1','2','3') else '1'
        questions = _mental_questions(grade); practice = _practice_questions(grade)
        practice_rows = ''.join(f'<tr><td>Ej. {i}</td><td>{escape(q)}</td><td class="answer">{escape(a)}</td></tr>' for i,(q,a) in enumerate(practice,1))
        question_rows = ''.join(f'<tr><td>{q["number"]}</td><td>{escape(q["prompt"])}</td><td class="answer">{escape(q["answer"])}</td><td class="category">{escape(q["category"])}</td></tr>' for q in questions)
        body = f'''<div class="no-print card"><h1>Generador de cálculo mental</h1><form method="get" class="grid"><label>Grado<select name="grade"><option value="1" {"selected" if grade=="1" else ""}>Primer grado</option><option value="2" {"selected" if grade=="2" else ""}>Segundo grado</option><option value="3" {"selected" if grade=="3" else ""}>Tercer grado</option></select></label><div><button>Generar otra prueba</button></div><div><button type="button" onclick="window.print()">Imprimir prueba</button></div><div><button type="button" onclick="document.body.classList.toggle('show-answers')">Mostrar u ocultar respuestas</button></div></form></div><div class="card mental-sheet"><h1>CÁLCULO MENTAL · {grade}.º DE SECUNDARIA</h1><p>Aplicación individual. Plantee cada pregunta de forma oral. Muestre apoyo visual únicamente conforme al procedimiento SiSAT.</p><table><tr><th>No.</th><th>Pregunta</th><th class="answer">Respuesta</th><th class="category">Habilidad equivalente</th></tr>{practice_rows}{question_rows}</table><p class="muted">Códigos de registro: 1 correcta sin apoyo visual · 1V correcta con apoyo visual · 0 equivocada o sin respuesta.</p></div><style>.mental-sheet{{max-width:920px;margin:auto}}.mental-sheet table{{min-width:0}}.mental-sheet td:first-child{{width:70px;text-align:center;font-weight:800}}.answer,.category{{display:none}}.show-answers .answer,.show-answers .category{{display:table-cell}}@media print{{.sidebar,.topbar,.bottom-nav,.footer,.workspace-banner,.workspace-mobile,.no-print{{display:none!important}}.main-area,.wrap{{margin:0!important;padding:0!important;max-width:none!important}}.mental-sheet{{box-shadow:none!important;border:0!important}}.mental-sheet .answer,.mental-sheet .category{{display:none!important}}}}</style>'''
        return core.page('Prueba de cálculo mental', body)

    @app.route('/sisat/export.xlsx')
    def sisat_export():
        if not session.get('uid'): return redirect('/login')
        if not _can_export(): flash('Solo Administración y Dirección pueden exportar concentrados SiSAT.'); return redirect('/sisat')
        scope = request.args.get('scope', 'group'); scope = scope if scope in ('school','group','student') else 'group'
        skill = request.args.get('skill', ''); student_id = request.args.get('student_id', type=int)
        output = _export_workbook(scope, skill, student_id)
        return send_file(output, as_attachment=True, download_name=f'sisat_{scope}_{date.today().isoformat()}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.after_request
    def sisat_navigation(response):
        if 'text/html' not in response.headers.get('Content-Type', '') or not session.get('uid'): return response
        html = response.get_data(as_text=True)
        if 'href="/sisat"' not in html:
            marker = '<a class="nav-link logout" href="/logout">'; link = '<a class="nav-link" href="/sisat"><span class="nav-icon">▦</span><span>SiSAT</span></a>'
            if marker in html: html = html.replace(marker, link + marker, 1)
        response.set_data(html); response.headers['Content-Length'] = str(len(response.get_data())); return response
