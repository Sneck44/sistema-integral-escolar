import io
import re
import unicodedata
from html import escape

import xlsxwriter
from flask import request, redirect, session, flash, send_file
from itsdangerous import URLSafeSerializer, BadSignature
from openpyxl import load_workbook

import app as core


MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 500


def _norm(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', ' ', text.lower()).strip()


def _text(value):
    return str(value or '').strip()


def _number(value, minimum=None, maximum=None):
    if value in (None, ''):
        return None
    try:
        result = float(str(value).replace(',', '.'))
    except (TypeError, ValueError):
        return None
    if minimum is not None and result < minimum:
        return None
    if maximum is not None and result > maximum:
        return None
    return result


def _integer(value):
    number = _number(value, 1)
    return int(number) if number is not None and number.is_integer() else None


def _split_name(full_name):
    parts = _text(full_name).split()
    if len(parts) >= 3:
        return parts[0], parts[1], ' '.join(parts[2:])
    if len(parts) == 2:
        return parts[0], '', parts[1]
    return '', '', parts[0] if parts else ''


def _read_sheet(upload):
    if not upload or not upload.filename:
        raise ValueError('Selecciona un archivo de Excel.')
    if not upload.filename.lower().endswith('.xlsx'):
        raise ValueError('El archivo debe estar en formato .xlsx.')
    raw = upload.read(MAX_FILE_SIZE + 1)
    if len(raw) > MAX_FILE_SIZE:
        raise ValueError('El archivo supera el límite de 5 MB.')
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    except Exception as exc:
        raise ValueError('No fue posible leer el archivo. Verifica que sea un Excel .xlsx válido.') from exc
    if not rows:
        raise ValueError('El archivo está vacío.')
    return rows


def _table(rows, required_aliases):
    header_index = None
    header_map = {}
    for index, row in enumerate(rows[:12]):
        candidate = {_norm(value): column for column, value in enumerate(row) if _norm(value)}
        if all(any(alias in candidate for alias in aliases) for aliases in required_aliases):
            header_index, header_map = index, candidate
            break
    if header_index is None:
        raise ValueError('No se reconocieron los encabezados. Descarga y utiliza la plantilla del sistema.')
    data = []
    for excel_row, row in enumerate(rows[header_index + 1:header_index + 1 + MAX_ROWS], start=header_index + 2):
        if any(value not in (None, '') for value in row):
            data.append((excel_row, row, header_map))
    return data


def _cell(row, headers, *aliases):
    for alias in aliases:
        column = headers.get(_norm(alias))
        if column is not None and column < len(row):
            return row[column]
    return None


def _student_index():
    by_number, by_name = {}, {}
    for student in core.Student.query.all():
        if student.list_no:
            by_number[student.list_no] = student
        by_name[_norm(student.full_name)] = student
    return by_number, by_name


def _student_duplicate(record, by_number, by_name):
    return by_number.get(record.get('list_no')) or by_name.get(_norm(
        f"{record.get('paternal', '')} {record.get('maternal', '')} {record.get('names', '')}"
    ))


def _parse_students(rows):
    aliases = [
        ('apellido paterno', 'alumno', 'nombre completo'),
        ('nombres', 'nombre s', 'alumno', 'nombre completo'),
    ]
    table = _table(rows, aliases)
    by_number, by_name = _student_index()
    records = []
    for excel_row, row, headers in table:
        full_name = _text(_cell(row, headers, 'Alumno', 'Nombre completo'))
        paternal = _text(_cell(row, headers, 'Apellido paterno'))
        maternal = _text(_cell(row, headers, 'Apellido materno'))
        names = _text(_cell(row, headers, 'Nombre(s)', 'Nombres'))
        if full_name and not (paternal or names):
            paternal, maternal, names = _split_name(full_name)
        record = {
            'row': excel_row,
            'list_no': _integer(_cell(row, headers, 'No.', 'No. de lista', 'Número de lista')),
            'paternal': paternal, 'maternal': maternal, 'names': names,
            'status': _text(_cell(row, headers, 'Estado')) or 'ACTIVO',
            'tutor': _text(_cell(row, headers, 'Tutor')),
            'phone': _text(_cell(row, headers, 'Teléfono', 'Telefono')),
            'weight_kg': _number(_cell(row, headers, 'Peso kg', 'Peso (kg)'), 0),
            'height_cm': _number(_cell(row, headers, 'Estatura cm', 'Estatura (cm)'), 0),
            'top_size': _text(_cell(row, headers, 'Playera/Blusa', 'Talla playera / blusa')),
            'bottom_size': _text(_cell(row, headers, 'Pantalón/Falda', 'Talla pantalón / falda')),
            'sweater_size': _text(_cell(row, headers, 'Suéter/Chamarra', 'Talla suéter / chamarra')),
            'shoe_size': _text(_cell(row, headers, 'Calzado', 'Número de calzado')),
            'uniform_notes': _text(_cell(row, headers, 'Observaciones')),
        }
        errors = []
        if not paternal: errors.append('Falta apellido paterno')
        if not names: errors.append('Falta nombre')
        duplicate = _student_duplicate(record, by_number, by_name) if not errors else None
        record['existing_id'] = duplicate.id if duplicate else None
        record['result'] = 'ERROR' if errors else ('DUPLICADO' if duplicate else 'NUEVO')
        record['errors'] = '; '.join(errors)
        records.append(record)
    return records


def _parse_diagnostic(rows):
    from diagnostic import SUBJECTS_BY_GRADE, _active_grade_key
    grade_key = _active_grade_key()
    subjects = SUBJECTS_BY_GRADE[grade_key]
    table = _table(rows, [('alumno', 'nombre completo'), ('no', 'no de lista', 'numero de lista')])
    by_number, by_name = _student_index()
    records = []
    for excel_row, row, headers in table:
        list_no = _integer(_cell(row, headers, 'No.', 'No. de lista', 'Número de lista'))
        full_name = _text(_cell(row, headers, 'Alumno', 'Nombre completo'))
        student = by_number.get(list_no) or by_name.get(_norm(full_name))
        scores = {}
        for subject in subjects:
            scores[subject] = _number(_cell(row, headers, subject), 0, 10)
        errors = []
        if not student: errors.append('Alumno no localizado en el grupo activo')
        invalid = []
        for subject in subjects:
            raw = _cell(row, headers, subject)
            if raw not in (None, '') and scores[subject] is None:
                invalid.append(subject)
        if invalid: errors.append('Calificación inválida en: ' + ', '.join(invalid))
        captured = [value for value in scores.values() if value is not None]
        overall = round(sum(captured) / len(captured), 2) if captured else _number(_cell(row, headers, 'Calificación', 'Promedio'), 0, 10)
        records.append({
            'row': excel_row, 'student_id': student.id if student else None,
            'list_no': list_no, 'student_name': student.full_name if student else full_name,
            'scores': scores, 'diagnostic_score': overall,
            'learning_pace': _text(_cell(row, headers, 'Ritmo')) or 'POR DETERMINAR',
            'learning_style': _text(_cell(row, headers, 'Estilo de aprendizaje', 'Estilo')).upper() or 'POR DETERMINAR',
            'perception_channel': _text(_cell(row, headers, 'Canal de percepción', 'Canal')).upper() or 'POR DETERMINAR',
            'performance_level': _text(_cell(row, headers, 'Nivel')).upper() or 'POR DETERMINAR',
            'strengths': _text(_cell(row, headers, 'Fortalezas')),
            'support_needs': _text(_cell(row, headers, 'Necesidades de apoyo')),
            'observations': _text(_cell(row, headers, 'Observaciones')),
            'result': 'ERROR' if errors else 'ACTUALIZAR', 'errors': '; '.join(errors),
            'grade_key': grade_key,
        })
    return records


def _signer():
    return URLSafeSerializer(core.app.secret_key, salt='excel-import-preview-v1')


def _preview(kind, records):
    valid = sum(1 for record in records if record['result'] != 'ERROR')
    errors = len(records) - valid
    rows = ''
    for record in records:
        label = record.get('student_name') or f"{record.get('paternal','')} {record.get('maternal','')} {record.get('names','')}".strip()
        badge = record['result']
        detail = record.get('errors') or ('Se actualizará el registro existente' if badge in ('DUPLICADO', 'ACTUALIZAR') else 'Se agregará como alumno nuevo')
        rows += f'<tr><td>{record["row"]}</td><td>{escape(str(record.get("list_no") or "—"))}</td><td><b>{escape(label)}</b></td><td><span class="import-badge {badge.lower()}">{escape(badge.title())}</span></td><td>{escape(detail)}</td></tr>'
    token = _signer().dumps({'kind': kind, 'records': records})
    mode = ''
    if kind == 'students':
        mode = '''<label style="display:flex;gap:8px;align-items:flex-start"><input style="width:auto;margin-top:4px" type="radio" name="mode" value="add_only" checked><span><b>Agregar solo alumnos nuevos</b><small style="display:block">Los duplicados no se modificarán.</small></span></label><label style="display:flex;gap:8px;align-items:flex-start"><input style="width:auto;margin-top:4px" type="radio" name="mode" value="update"><span><b>Agregar nuevos y actualizar existentes</b><small style="display:block">Completa o reemplaza los datos del alumno coincidente.</small></span></label>'''
    return core.page('Vista previa de importación', f'''
      <h1>Vista previa · {'Alumnos' if kind == 'students' else 'Diagnóstico'}</h1>
      <div class="grid"><div class="card"><div class="muted">Filas leídas</div><div class="kpi">{len(records)}</div></div><div class="card"><div class="muted">Listas para importar</div><div class="kpi">{valid}</div></div><div class="card"><div class="muted">Con errores</div><div class="kpi">{errors}</div></div></div>
      <div class="card"><h2>Revisa antes de guardar</h2><p class="muted">Las filas con error se omitirán. La información todavía no se ha guardado.</p><div class="scroll"><table><tr><th>Fila</th><th>No.</th><th>Alumno</th><th>Resultado</th><th>Detalle</th></tr>{rows}</table></div></div>
      <form method="post" action="/imports/{kind}/confirm" class="card"><input type="hidden" name="token" value="{escape(token)}"><div class="grid">{mode}</div><div style="display:flex;gap:10px;flex-wrap:wrap;margin-top:16px"><button style="width:auto" {'disabled' if not valid else ''}>Confirmar importación</button><a href="/{'students' if kind == 'students' else 'diagnostic'}" style="padding:10px 14px">Cancelar</a></div></form>
      <style>.import-badge{{display:inline-block;padding:5px 9px;border-radius:99px;font-size:11px;font-weight:850}}.import-badge.nuevo{{background:#e4f6e9;color:#176b35}}.import-badge.duplicado,.import-badge.actualizar{{background:#fff2cf;color:#78580d}}.import-badge.error{{background:#fde4e4;color:#982626}}</style>''')


def _template(kind):
    output = io.BytesIO()
    wb = xlsxwriter.Workbook(output, {'in_memory': True})
    ws = wb.add_worksheet('Alumnos' if kind == 'students' else 'Diagnóstico')
    title = wb.add_format({'bold': True, 'font_size': 16, 'font_color': '#FFFFFF', 'bg_color': '#7B1024', 'align': 'center', 'valign': 'vcenter'})
    header = wb.add_format({'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#217346', 'text_wrap': True, 'align': 'center', 'valign': 'vcenter', 'border': 1})
    note = wb.add_format({'font_color': '#6B5960', 'bg_color': '#FFF8ED', 'text_wrap': True})
    if kind == 'students':
        headers = ['No. de lista', 'Apellido paterno', 'Apellido materno', 'Nombre(s)', 'Estado', 'Tutor', 'Teléfono', 'Peso (kg)', 'Estatura (cm)', 'Talla playera / blusa', 'Talla pantalón / falda', 'Talla suéter / chamarra', 'Número de calzado', 'Observaciones']
        title_text = 'PLANTILLA PARA IMPORTAR ALUMNOS'
        example = [1, 'PÉREZ', 'GARCÍA', 'ANA SOFÍA', 'ACTIVO', 'MARÍA GARCÍA', '7971234567', 45.5, 152, 'M', '14', 'M', '24.5', '']
    else:
        from diagnostic import _subjects_for_active_grade
        headers = ['No. de lista', 'Alumno'] + _subjects_for_active_grade() + ['Estilo de aprendizaje', 'Canal de percepción', 'Nivel', 'Fortalezas', 'Necesidades de apoyo', 'Observaciones']
        title_text = 'PLANTILLA PARA IMPORTAR EVALUACIÓN DIAGNÓSTICA'
        example = [1, 'PÉREZ GARCÍA ANA SOFÍA'] + [8] * len(_subjects_for_active_grade()) + ['DIVERGENTE', 'VISUAL', 'ESPERADO', 'Participa y argumenta', 'Reforzar cálculo mental', '']
    ws.merge_range(0, 0, 0, len(headers) - 1, title_text, title)
    ws.merge_range(1, 0, 1, len(headers) - 1, 'No cambies los encabezados. Puedes borrar la fila de ejemplo antes de importar.', note)
    for col, value in enumerate(headers): ws.write(3, col, value, header)
    for col, value in enumerate(example): ws.write(4, col, value)
    ws.freeze_panes(4, 0); ws.autofilter(3, 0, 4, len(headers) - 1)
    ws.set_column(0, 0, 12); ws.set_column(1, len(headers) - 1, 20)
    wb.close(); output.seek(0)
    return output


def _save_students(records, mode):
    from student_details import StudentDetails
    by_number, by_name = _student_index()
    created = updated = skipped = 0
    for record in records:
        if record['result'] == 'ERROR': skipped += 1; continue
        student = _student_duplicate(record, by_number, by_name)
        if student and mode != 'update': skipped += 1; continue
        if not student:
            student = core.Student(**{
                field: record.get(field)
                for field in ('list_no', 'paternal', 'maternal', 'names', 'status', 'tutor', 'phone')
            })
            core.db.session.add(student); core.db.session.flush(); created += 1
        else:
            updated += 1
            for field in ('list_no', 'paternal', 'maternal', 'names', 'status', 'tutor', 'phone'):
                setattr(student, field, record.get(field))
        details = StudentDetails.query.filter_by(student_id=student.id).first()
        if not details: details = StudentDetails(student_id=student.id); core.db.session.add(details)
        for field in ('weight_kg', 'height_cm', 'top_size', 'bottom_size', 'sweater_size', 'shoe_size', 'uniform_notes'):
            setattr(details, field, record.get(field))
        if student.list_no: by_number[student.list_no] = student
        by_name[_norm(student.full_name)] = student
    core.db.session.commit()
    return created, updated, skipped


def _save_diagnostic(records):
    from diagnostic import StudentDiagnostic, DiagnosticSubjectGrade
    updated = skipped = 0
    for record in records:
        if record['result'] == 'ERROR': skipped += 1; continue
        d = StudentDiagnostic.query.filter_by(student_id=record['student_id']).first()
        if not d: d = StudentDiagnostic(student_id=record['student_id']); core.db.session.add(d)
        for field in ('diagnostic_score', 'learning_pace', 'learning_style', 'perception_channel', 'performance_level', 'strengths', 'support_needs', 'observations'):
            setattr(d, field, record.get(field))
        for subject, score in record['scores'].items():
            if score is None: continue
            grade = DiagnosticSubjectGrade.query.filter_by(student_id=record['student_id'], grade_level=record['grade_key'], subject_name=subject).first()
            if not grade: grade = DiagnosticSubjectGrade(student_id=record['student_id'], grade_level=record['grade_key'], subject_name=subject); core.db.session.add(grade)
            grade.score = score
        updated += 1
    core.db.session.commit()
    return updated, skipped


def install(app):
    @app.route('/imports/students', methods=['POST'])
    def import_students():
        if not session.get('uid'): return redirect('/login')
        try: return _preview('students', _parse_students(_read_sheet(request.files.get('file'))))
        except ValueError as exc: flash(str(exc)); return redirect('/students')

    @app.route('/imports/diagnostic', methods=['POST'])
    def import_diagnostic():
        if not session.get('uid'): return redirect('/login')
        try: return _preview('diagnostic', _parse_diagnostic(_read_sheet(request.files.get('file'))))
        except ValueError as exc: flash(str(exc)); return redirect('/diagnostic')

    @app.route('/imports/<kind>/confirm', methods=['POST'])
    def confirm_import(kind):
        if not session.get('uid'): return redirect('/login')
        try: payload = _signer().loads(request.form.get('token', ''))
        except BadSignature: flash('La vista previa expiró o fue modificada. Vuelve a cargar el archivo.'); return redirect('/students' if kind == 'students' else '/diagnostic')
        if payload.get('kind') != kind or kind not in ('students', 'diagnostic'): return redirect('/')
        if kind == 'students':
            created, updated, skipped = _save_students(payload['records'], request.form.get('mode', 'add_only'))
            flash(f'Importación terminada: {created} alumnos nuevos, {updated} actualizados y {skipped} omitidos.')
            return redirect('/students')
        updated, skipped = _save_diagnostic(payload['records'])
        flash(f'Importación terminada: {updated} diagnósticos actualizados y {skipped} filas omitidas.')
        return redirect('/diagnostic')

    @app.route('/templates/<kind>.xlsx')
    def download_template(kind):
        if not session.get('uid'): return redirect('/login')
        if kind not in ('students', 'diagnostic'): return redirect('/')
        return send_file(_template(kind), as_attachment=True, download_name=f'plantilla_{kind}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.after_request
    def import_buttons(response):
        if 'text/html' not in response.headers.get('Content-Type', '') or not session.get('uid'): return response
        html = response.get_data(as_text=True)
        if request.path == '/students' and 'action="/imports/students"' not in html:
            block = '''<div class="card"><h2>Importar o exportar lista de alumnos</h2><p class="muted">Carga varios alumnos desde Excel con vista previa y detección de duplicados.</p><form method="post" action="/imports/students" enctype="multipart/form-data" style="display:flex;gap:10px;align-items:end;flex-wrap:wrap"><label style="flex:1;min-width:240px">Archivo Excel (.xlsx)<input type="file" name="file" accept=".xlsx" required></label><button style="width:auto">Revisar e importar</button><a href="/templates/students.xlsx" style="padding:10px 14px;font-weight:700">Descargar plantilla</a><a href="/exports/students.xlsx" style="padding:10px 14px;font-weight:700">Exportar lista</a></form></div>'''
            html = html.replace('<div class="card scroll">', block + '<div class="card scroll">', 1)
        if request.path == '/diagnostic' and 'action="/imports/diagnostic"' not in html:
            block = '''<div class="card"><h2>Importar evaluación diagnóstica desde Excel</h2><p class="muted">Relaciona cada fila con el alumno del grupo activo y permite revisar los resultados antes de guardarlos.</p><form method="post" action="/imports/diagnostic" enctype="multipart/form-data" style="display:flex;gap:10px;align-items:end;flex-wrap:wrap"><label style="flex:1;min-width:240px">Archivo Excel (.xlsx)<input type="file" name="file" accept=".xlsx" required></label><button style="width:auto">Revisar e importar</button><a href="/templates/diagnostic.xlsx" style="padding:10px 14px;font-weight:700">Descargar plantilla</a></form></div>'''
            html = html.replace('<div class="card">\n          <h2>Registro diagnóstico por asignatura</h2>', block + '<div class="card">\n          <h2>Registro diagnóstico por asignatura</h2>', 1)
        response.set_data(html); response.headers['Content-Length'] = str(len(response.get_data())); return response
